"""SummarizeAI Pro — local-first PDF workspace. Run: py -3.11 -m streamlit run pdfsummarizer_pro.py"""
import io
import re
import hashlib
from collections import Counter
from html import escape

import streamlit as st
import pandas as pd
import plotly.express as px
from pypdf import PdfReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document

st.set_page_config(page_title="SummarizeAI Pro", page_icon="🧠", layout="wide")
st.markdown("""<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Space+Grotesk:wght@500;600;700&display=swap');
html,body,[class*="css"],.stApp {font-family:'DM Sans',sans-serif}
.stApp {background:radial-gradient(ellipse at 78% 0%,#25265a 0%,transparent 36%),radial-gradient(ellipse at 0% 60%,#102b43 0%,transparent 42%),#090e1d;color:#e8edff}
[data-testid="stSidebar"] {background:linear-gradient(180deg,#141a35,#0c1329);border-right:1px solid #30395a}
[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p {color:#bfc9e4}
.block-container {padding-top:2rem;max-width:1350px;padding-bottom:4rem}
h1,h2,h3 {font-family:'Space Grotesk',sans-serif!important;color:#f4f4ff!important;letter-spacing:-.025em}
h2 {margin-top:1.3rem!important}
[data-testid="stMetric"] {background:linear-gradient(145deg,rgba(39,48,86,.95),rgba(20,29,54,.95));border:1px solid #39466e;border-radius:18px;padding:19px;box-shadow:0 14px 30px #05091644}
[data-testid="stMetricValue"] {color:#fff;font-family:'Space Grotesk',sans-serif}
[data-testid="stMetricLabel"] {color:#b9c5e4}
.stButton>button[kind="primary"],.stDownloadButton>button[kind="primary"] {background:linear-gradient(100deg,#7358f5,#467cfa);border:0;color:white;border-radius:12px;box-shadow:0 8px 22px #5847d744}
.stButton>button,.stDownloadButton>button {border-radius:11px;transition:transform .18s ease,box-shadow .18s ease}
.stButton>button:hover,.stDownloadButton>button:hover {transform:translateY(-2px);box-shadow:0 7px 20px #514cc433}
[data-testid="stFileUploader"] {background:#151e39;border:1.5px dashed #6d76ba;border-radius:18px;padding:18px}
[data-testid="stTabs"] [data-baseweb="tab-list"] {gap:12px;background:#111a31;padding:9px;border:1px solid #303c60;border-radius:14px}
[data-testid="stTabs"] button {border-radius:9px;color:#bfc8e3;padding:10px 16px}
[data-testid="stTabs"] button[aria-selected="true"] {background:#5146a8;color:white}
[data-testid="stTabs"] [data-baseweb="tab-highlight"] {display:none}
[data-testid="stTextInput"] input,[data-testid="stTextArea"] textarea {background:#141d36;color:#f0f2ff;border-radius:10px}
.hero {position:relative;overflow:hidden;padding:43px 40px;border:1px solid #535b91;border-radius:24px;background:linear-gradient(112deg,#30245f 0%,#182b59 58%,#102a41 100%);box-shadow:0 22px 60px #03081670}
.hero:after {content:'✦';position:absolute;right:8%;top:5%;font-size:125px;color:#ffffff16;transform:rotate(-12deg)}
.hero h1 {font-size:clamp(2.3rem,5vw,4.2rem);line-height:1.08;margin:10px 0 13px}
.hero h1 span {color:#b7a7ff}
.hero p {color:#c9d4ef;font-size:1.05rem;max-width:670px;line-height:1.6}
.eyebrow {color:#c4b5fd;letter-spacing:.18em;font-weight:700;font-size:.75rem}
.hero-chips {display:flex;gap:9px;flex-wrap:wrap;margin-top:22px}
.hero-chips span {background:#ffffff13;border:1px solid #ffffff35;border-radius:99px;padding:7px 12px;color:#e8eaff;font-size:.8rem}
.section-kicker {color:#9fa9cc;font-size:.8rem;letter-spacing:.12em;font-weight:700;margin:26px 0 6px}
</style>""", unsafe_allow_html=True)

STOP = set("a an and are as at be been by for from has have in into is it its of on or that the their this to was were will with you your can not but than then also such these those".split())
CATEGORIES = {
    "Research paper": ["abstract", "methodology", "literature review", "references", "hypothesis", "experiment"],
    "Resume / CV": ["work experience", "education", "skills", "curriculum vitae", "professional experience"],
    "Invoice": ["invoice", "bill to", "subtotal", "tax", "amount due", "payment terms"],
    "Study notes": ["chapter", "learning objectives", "exercise", "lecture", "revision", "questions"],
    "Business report": ["executive summary", "revenue", "quarter", "stakeholders", "recommendations", "performance"],
    "Legal document": ["agreement", "hereby", "terms and conditions", "party", "jurisdiction", "liability"],
}

def classify(text):
    sample = text[:30000].lower()
    scores = {name:sum(sample.count(term) for term in terms) for name,terms in CATEGORIES.items()}
    name = max(scores,key=scores.get)
    return (name if scores[name] else "General document", scores[name])

def keywords(text,n=15):
    tokens = re.findall(r"[A-Za-z][A-Za-z'-]{2,}",text.lower())
    return Counter(w for w in tokens if w not in STOP).most_common(n)

@st.cache_data(show_spinner=False)
def extract_pdf(data, use_ocr, ocr_lang):
    reader=PdfReader(io.BytesIO(data))
    pages=[]; ocr_pages=[]
    for i,page in enumerate(reader.pages):
        try: value=page.extract_text() or ""
        except Exception: value=""
        if use_ocr and len(value.strip())<35:
            try:
                import fitz, pytesseract
                from PIL import Image
                with fitz.open(stream=data,filetype="pdf") as doc:
                    pix=doc[i].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
                    image=Image.open(io.BytesIO(pix.tobytes("png")))
                    value=pytesseract.image_to_string(image,lang=ocr_lang)
                    ocr_pages.append(i+1)
            except Exception as exc:
                raise RuntimeError(f"OCR failed on page {i+1}. Install Tesseract OCR and its language data. Details: {exc}") from exc
        pages.append(value.strip())
    return pages,ocr_pages

@st.cache_resource(show_spinner=False)
def get_summarizer(offline):
    from transformers import pipeline
    model="sshleifer/distilbart-cnn-12-6"
    if offline:
        # Only use already-cached model files; no network request.
        from transformers import AutoTokenizer, AutoModelForSeq2SeqLM
        tokenizer=AutoTokenizer.from_pretrained(model,local_files_only=True)
        network=AutoModelForSeq2SeqLM.from_pretrained(model,local_files_only=True)
        return pipeline("summarization",model=network,tokenizer=tokenizer,device=-1)
    return pipeline("summarization",model=model,device=-1)

def summarize(text,mode,offline):
    model=get_summarizer(offline)
    words=text.split()
    chunks=[" ".join(words[i:i+360]) for i in range(0,len(words),360)]
    limits={"Short":(75,20),"Medium":(120,35),"Detailed":(175,50)}
    maximum,minimum=limits[mode]
    outputs=[]
    bar=st.progress(0,text="Summarizing locally...")
    for i,chunk in enumerate(chunks):
        # Skip generation for very short text; avoid model's short-input warnings.
        if len(chunk.split())<45: outputs.append(chunk)
        else:
            result=model(chunk,max_length=maximum,min_length=minimum,do_sample=False,truncation=True)
            outputs.append(result[0]["summary_text"].strip())
        bar.progress((i+1)/len(chunks),text=f"Processed {i+1}/{len(chunks)} sections")
    bar.empty()
    return "\n\n".join(outputs),len(chunks)

def report_content(name,summary,takeaways,stats):
    return {"name":name,"summary":summary,"takeaways":takeaways,"stats":stats}

def export_pdf(report):
    out=io.BytesIO(); doc=SimpleDocTemplate(out)
    styles=getSampleStyleSheet()
    story=[Paragraph("SummarizeAI Pro — Document Report",styles["Title"]),Spacer(1,14),
           Paragraph(escape(report["name"]),styles["Heading2"]),
           Paragraph(escape(report["stats"]),styles["Normal"]),Spacer(1,14),
           Paragraph("AI Summary",styles["Heading2"])]
    for paragraph in report["summary"].split("\n\n"):
        story.append(Paragraph(escape(paragraph).replace("\n","<br/>"),styles["BodyText"]))
        story.append(Spacer(1,8))
    story.append(Paragraph("Key Takeaways",styles["Heading2"]))
    for point in report["takeaways"]: story.append(Paragraph("- "+escape(point),styles["BodyText"]))
    doc.build(story); return out.getvalue()

def export_docx(report):
    doc=Document(); doc.add_heading("SummarizeAI Pro — Document Report",0)
    doc.add_heading(report["name"],level=1); doc.add_paragraph(report["stats"])
    doc.add_heading("AI Summary",level=1)
    for paragraph in report["summary"].split("\n\n"):doc.add_paragraph(paragraph)
    doc.add_heading("Key Takeaways",level=1)
    for point in report["takeaways"]:doc.add_paragraph(point,style="List Bullet")
    out=io.BytesIO();doc.save(out);return out.getvalue()

st.markdown('''<div class="hero"><div class="eyebrow">✦ YOUR INTELLIGENT DOCUMENT WORKSPACE</div><h1>Meet <span>SummarizeAI Pro.</span><br>Read less. Understand more.</h1><p>Turn lengthy documents into clear insights with local AI summarization, OCR, powerful search and interactive analytics.</p><div class="hero-chips"><span>✦ AI-powered summaries</span><span>◈ Scanned PDF OCR</span><span>⌕ Smart search</span><span>▥ Analytics</span></div></div>''',unsafe_allow_html=True)
st.write("")
with st.sidebar:
    st.markdown("### ◈ SummarizeAI Pro")
    st.caption("DOCUMENT INTELLIGENCE STUDIO")
    st.divider()
    st.header("⚙️ Workspace Settings")
    summary_mode=st.selectbox("Summary length",["Short","Medium","Detailed"],index=1)
    use_ocr=st.toggle("OCR for scanned pages",value=True)
    ocr_lang=st.text_input("Tesseract language code",value="eng",help="Use eng for English. Other languages require installed Tesseract language packs.")
    offline=st.toggle("Strict offline model mode",value=False,help="Requires the model to have been downloaded previously. No external inference API is used in either mode.")
    st.caption("Your uploaded PDF is processed in this Streamlit process. Initial AI model download may require internet unless strict offline mode is enabled.")

st.markdown('<div class="section-kicker">01 / GET STARTED</div>',unsafe_allow_html=True)
st.subheader("Your document, transformed")
st.caption("Upload a PDF below to unlock your AI workspace.")
uploaded=st.file_uploader("Upload your PDF document",type=["pdf"])
if not uploaded:
    st.info("Upload a PDF to unlock document intelligence, searchable pages, bookmarks and AI summary export.")
    st.stop()
raw=uploaded.getvalue(); document_id=hashlib.sha256(raw).hexdigest()
if st.session_state.get("document_id")!=document_id:
    st.session_state.document_id=document_id
    st.session_state.bookmarks=[]
    st.session_state.summary=""
    st.session_state.summary_chunks=0
    st.session_state.summary_mode_used=""

try:
    with st.spinner("Extracting document text (and running OCR where needed)..."):
        pages,ocr_pages=extract_pdf(raw,use_ocr,ocr_lang)
except Exception as exc:
    st.error(str(exc));st.stop()
text="\n\n".join(f"[Page {i+1}]\n{p}" for i,p in enumerate(pages) if p)
if not text.strip():
    st.error("No extractable text found. Enable OCR in the sidebar and ensure Tesseract OCR is installed.");st.stop()
wc=[len(p.split()) for p in pages]; total=sum(wc); label,score=classify(text)
kw=keywords(text)

c1,c2,c3,c4=st.columns(4)
c1.metric("Pages",len(pages));c2.metric("Words",f"{total:,}")
c3.metric("Reading time",f"{max(1,round(total/225))} min");c4.metric("OCR pages",len(ocr_pages))
st.caption(f"Detected document type: **{label}** (heuristic estimate, not a verified classification) · File size: {len(raw)/1024/1024:.2f} MB")

tab1,tab2,tab3,tab4=st.tabs(["✨ AI Summary","🔎 Document Explorer","📊 Analytics","🔖 Bookmarks"])
with tab1:
    st.subheader("AI Summary Studio")
    if st.button("✨ Generate AI Summary",type="primary",use_container_width=True):
        try:
            with st.spinner("Loading the local AI model..."):
                summary,chunks=summarize("\n\n".join(pages),summary_mode,offline)
            st.session_state.summary=summary
            st.session_state.summary_chunks=chunks
            st.session_state.summary_mode_used=summary_mode
        except Exception as exc:
            st.error(f"Summarization failed: {exc}")
            if offline: st.info("Strict offline mode requires the model files to have been downloaded and cached previously. Disable it for the first download.")
    if st.session_state.summary:
        summary=st.session_state.summary
        st.success(f"Summary ready · {st.session_state.summary_mode_used} · {st.session_state.summary_chunks} sections")
        st.write(summary)
        st.caption("Use the copy icon in the top-right of the text block below for one-click copying.")
        st.code(summary,language=None)
        points=[s.strip() for s in re.split(r"(?<=[.!?])\s+",summary.replace("\n"," ")) if len(s.strip())>30][:5]
        st.subheader("Key takeaways")
        for point in points:st.markdown("- "+point)
        report=report_content(uploaded.name,summary,points,f"Pages: {len(pages)} | Words: {total} | Category: {label}")
        a,b,c=st.columns(3)
        a.download_button("⬇ TXT",summary,file_name="AI_Summary.txt",mime="text/plain",use_container_width=True)
        b.download_button("⬇ PDF",export_pdf(report),file_name="AI_Summary.pdf",mime="application/pdf",use_container_width=True)
        c.download_button("⬇ Word",export_docx(report),file_name="AI_Summary.docx",mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True)
    else: st.info("Choose your summary length in the sidebar, then generate a summary.")

with tab2:
    st.subheader("Search your PDF")
    query=st.text_input("Search for a word or phrase",placeholder="e.g. methodology")
    if query.strip():
        hits=[]
        for i,p in enumerate(pages):
            for match in list(re.finditer(re.escape(query.strip()),p,re.I))[:30]:
                snippet=p[max(0,match.start()-110):min(len(p),match.end()+130)].replace("\n"," ")
                hits.append((i+1,snippet))
        st.caption(f"{len(hits)} matches shown (up to 30 per page)")
        for j,(page,snippet) in enumerate(hits):
            st.markdown(f"**Page {page}** — {escape(snippet)}")
            if st.button("🔖 Bookmark result",key=f"hit-{j}"):
                entry=(page,snippet)
                if entry not in st.session_state.bookmarks:st.session_state.bookmarks.append(entry)
                st.success("Bookmarked")
    st.subheader("Page explorer")
    chosen=st.selectbox("Select page",list(range(1,len(pages)+1)))
    st.text_area("Extracted page text",pages[chosen-1],height=300)
    if st.button("🔖 Bookmark this page"):
        entry=(chosen,pages[chosen-1][:220])
        if entry not in st.session_state.bookmarks:st.session_state.bookmarks.append(entry)
        st.success("Bookmarked")

with tab3:
    st.subheader("Interactive document analytics")
    df=pd.DataFrame({"Page":list(range(1,len(pages)+1)),"Words":wc})
    st.plotly_chart(px.bar(df,x="Page",y="Words",title="Words per page",template="plotly_dark"),use_container_width=True)
    if kw:
        kdf=pd.DataFrame(kw,columns=["Keyword","Frequency"])
        st.plotly_chart(px.bar(kdf,x="Frequency",y="Keyword",orientation="h",title="Top document keywords",template="plotly_dark"),use_container_width=True)
    st.dataframe(df,use_container_width=True,hide_index=True)
    st.caption("Keyword frequencies are statistical counts, not AI-generated topic labels.")

with tab4:
    st.subheader("Saved bookmarks")
    st.caption("Bookmarks remain available during this browser session for the current PDF. They are not saved permanently to disk.")
    if not st.session_state.bookmarks:st.info("Use Document Explorer to bookmark a passage or page.")
    for i,(page,snippet) in enumerate(list(st.session_state.bookmarks)):
        left,right=st.columns([5,1])
        left.markdown(f"**Page {page}:** {escape(snippet)}")
        if right.button("Remove",key=f"remove-{i}"):
            st.session_state.bookmarks.pop(i);st.rerun()
    if st.session_state.bookmarks:
        bookmarks_txt="\n\n".join(f"Page {p}: {s}" for p,s in st.session_state.bookmarks)
        st.download_button("Download bookmarks",bookmarks_txt,file_name="bookmarks.txt")

st.divider()
st.caption("SummarizeAI Pro · Python + Streamlit + PyPDF + Hugging Face · Local-first academic project")
